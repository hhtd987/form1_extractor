import csv
import json
import os
import re
from io import StringIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

os.chdir(r"C:/Users/pc/OneDrive/Desktop/usb_backup_29_7_2026/mycodes/pyprojects/form1_info_extract")

# =====================
# READ SQL FILE
# =====================
with open("responses_rows.sql", "r", encoding="utf-8") as f:
    sql = f.read()

# =====================
# ISOLATE VALUES PART
# =====================
values_part = sql.split("VALUES", 1)[1].strip().rstrip(";")

# =====================
# SPLIT ROWS CORRECTLY (FIXED ✅)
# =====================
raw_rows = re.findall(r"\((.*?)\)(?=,\s*\(|$)", values_part, re.DOTALL)

print(f"✅ RAW ROWS FOUND: {len(raw_rows)}")

# =====================
# PARSE EACH ROW (FIXED ✅)
# =====================

records = []

def parse_sql_row(row):
    fields = []
    current = ""
    in_string = False
    escape = False

    for c in row:
        if c == "\\" and in_string:
            escape = True
            current += c
            continue

        if c == "'" and not escape:
            in_string = not in_string
            continue

        if c == "," and not in_string:
            fields.append(current.strip())
            current = ""
        else:
            current += c

        escape = False

    if current:
        fields.append(current.strip())

    return fields


for row in raw_rows:
    try:
        fields = parse_sql_row(row)

        if len(fields) != 7:
            print(f"⚠ Skipping malformed row (fields={len(fields)})")
            continue

        eval_answers = json.loads(fields[4])
        text_answers = json.loads(fields[5])

        record = {
            "id": fields[0],
            "form_id": fields[1],
            "job_title": fields[2],
            "nature_of_work": fields[3],
            "submitted_at": fields[6],
        }

        # q_0 → q_23
        for i in range(24):
            record[f"q_{i}"] = eval_answers.get(str(i))

        # text answers
        for k, v in text_answers.items():
            record[f"text_{k}"] = v

        records.append(record)

    except Exception as e:
        print(f"⚠ Error parsing row: {e}")

# =====================
# ADD ORDERED NUMBER COLUMN
# =====================
for idx, record in enumerate(records, start=1):
    record["no"] = idx

# =====================
# FILL EMPTY CELLS
# =====================
NULL_AR = "لايوجد"

for record in records:
    for key, value in record.items():
        if value is None or (isinstance(value, str) and value.strip() == ""):
            record[key] = NULL_AR

# =====================
# VALIDATION
# =====================
print(f"✅ Total responses extracted: {len(records)}")

# =====================
# HEADERS
# =====================
fieldnames = ["no"] + [k for k in records[0].keys() if k != "no"]

group_header = [
    "المحور", "ID", "Form ID",
    "معلومات عامة", "معلومات عامة", "معلومات عامة",
    "الأهداف", "الأهداف",
    "تقييم البرنامج التدريبي", "تقييم البرنامج التدريبي",
    "تقييم البرنامج التدريبي", "تقييم البرنامج التدريبي",
    "المادة التدريبية", "المادة التدريبية",
    "المادة التدريبية", "المادة التدريبية",
    "تقييم المدرب", "تقييم المدرب", "تقييم المدرب",
    "تقييم المدرب", "تقييم المدرب", "تقييم المدرب",
    "التنسيق والإشراف العلمي", "التنسيق والإشراف العلمي",
    "التنسيق والإشراف العلمي",
    "مكان وزمان البرنامج التدريبي",
    "مكان وزمان البرنامج التدريبي",
    "مكان وزمان البرنامج التدريبي",
    "الخدمات الإدارية", "الخدمات الإدارية",
    "أسئلة أخرى", "أسئلة أخرى", "أسئلة أخرى", "أسئلة أخرى"
]

arabic_headers = [
    "ت", "ID", "Form ID",
    "العنوان الوظيفي",
    "طبيعة العمل الذي تمارسه",
    "تاريخ تقديم الطلب",
    "هل تجد نفسك قادراً على تعريف مفاهيم الذكاء الإصطناعي الأساسية؟",
    "هل تجد نفسك قادرا على فهم أنواع وتقنيات الذكاء الإصطناعي؟",
    "محتوى البرنامج التدريبي",
    "درجة تطوير البرنامج لمعارفك ومهاراتك",
    "ملائمة البرنامج لعملك",
    "إمكانية تطبيق ماتعلمته على عملك",
    "مدى ارتباط المادة التدريبية بموضوعات البرنامج",
    "مدى وضوح المادة التدريبية وتنظيمها وخلوها من الأخطاء",
    "مدى احتواء المادة التدريبية على موضوعات جديدة",
    "المادة التدريبية التي وزعت في البرنامج",
    "الأداء الحركي والصوتي للمدرب",
    "قدرة المدرب على توصيل المعلومات",
    "تنوع أساليب العرض من حيث الوسائل المستخدمة",
    "السيطرة على الوقت",
    "مدى تعاونه مع المتدربين",
    "قدرة المدرب على تحفيز المشاركين على التفاعل",
    "حضور المنسق كافة الجلسات التدريبية",
    "استماع المنسق لملاحظات واقتراحات المشاركين",
    "مدى تلبيته لمتطلبات واحتياجات المشاركين",
    "كفاية عدد الساعات التدريبية للموضوعات",
    "مدى ملائمة نظام الجلسات التدريبية",
    "مدى ملائمة توقيت البرنامج",
    "ملائمة القاعة التدريبية",
    "الظروف البيئية المحيطة بالمشارك (التهوية , الإضاءة , التكييف)",
    "الموضوعات غير المناسبة",
    "أهم المعارف والمهارات المكتسبة",
    "عنوان دورة مقترحة",
    "الملاحظات"
]

# =====================
# EXPORT XLSX
# =====================
output_file = "form_responses.xlsx"

wb = Workbook()
ws = wb.active
ws.title = "Form Responses"

# GROUP HEADER (ROW 1)
for col_idx, header in enumerate(group_header, start=1):
    ws.cell(row=1, column=col_idx, value=header)

# MERGE CELLS
start_col = 1
current_value = group_header[0]

for col_idx in range(2, len(group_header) + 2):
    if col_idx <= len(group_header) and group_header[col_idx - 1] == current_value:
        continue
    else:
        end_col = col_idx - 1
        if start_col != end_col:
            ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=end_col)
        start_col = col_idx
        if col_idx <= len(group_header):
            current_value = group_header[col_idx - 1]

# STYLE HEADER
for col in range(1, len(group_header) + 1):
    cell = ws.cell(row=1, column=col)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.font = Font(bold=True)

# ARABIC HEADERS (ROW 2)
for col_idx, header in enumerate(arabic_headers, start=1):
    cell = ws.cell(row=2, column=col_idx, value=header)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.font = Font(bold=True)

# DATA
for row_idx, record in enumerate(records, start=3):
    for col_idx, key in enumerate(fieldnames, start=1):
        ws.cell(row=row_idx, column=col_idx, value=record.get(key))

# COLUMN WIDTH
for col_idx in range(1, len(fieldnames) + 1):
    ws.column_dimensions[get_column_letter(col_idx)].width = 25

# FREEZE
ws.freeze_panes = "A3"

# SAVE
wb.save(output_file)

print(f"✅ XLSX created successfully: {output_file}")