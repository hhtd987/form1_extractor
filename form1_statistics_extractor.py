from pathlib import Path
import re
import json
from collections import Counter, defaultdict

# Arabic question headers (MCQs are from index 5 to 28)
arabic_headers = [
    "ID",
    "Form ID",
    "العنوان الوظيفي",
    "طبيعة العمل الذي تمارسه",
    "تاريخ تقديم الطلب",
    "هل تجد نفسك قادراً على إنشاء موقع ألكتروني شخصي؟",
    "هل تجد نفسك قادرا على تحديث ونشر وإجراء التديلات على المواقع الألكترونية؟",
    "محتوى البرنامج التدريبي",
    "درجة تطوير البرنامج لمعارفك ومهاراتك",
    "ملائمة البرنامج لعملك",
    "إمكانية تطبيق ماتعلمته على عملك",
    "مدى ارتباط المادة التدريبية بموضوعات البرنامج",
    "مدى وضوح المادة التدريبية وتنظيمها وخلوها من الأخطاء",
    "مدى احتواء المادة التدريبية على موضوعات جديدة",
    "المادة التدريبية التي وزعت في البرنامج",
    "الأداء الحركي والصوتي للمدرب",
    "قدرة المدرب على توصيل المعلومات",
    "تنوع أساليب العرض من حيث الوسائل المستخدمة",
    "السيطرة على الوقت",
    "مدى تعاونه مع المتدربين",
    "قدرة المدرب على تحفيز المشاركين على التفاعل",
    "حضور المنسق كافة الجلسات التدريبية",
    "استماع المنسق لملاحظات واقتراحات المشاركين",
    "مدى تلبيته لمتطلبات واحتياجات المشاركين",
    "كفاية عدد الساعات التدريبية للموضوعات",
    "مدى ملائمة نظام الجلسات التدريبية",
    "مدى ملائمة توقيت البرنامج",
    "ملائمة القاعة التدريبية",
    "الظروف البيئية المحيطة بالمشارك (التهوية , الإضاءة , التكييف)",
    "الموضوعات غير المناسبة",
    "أهم المعارف والمهارات المكتسبة",
    "عنوان دورة مقترحة",
    "الملاحظات"
]



BASE_DIR = Path(__file__).resolve().parent
SQL_PATH = BASE_DIR / "responses_rows.sql"
OUTPUT_PATH = BASE_DIR / "form_statistics.xlsx"


def split_sql_fields(row_text):
    """Split one SQL VALUES row on commas outside quoted strings."""
    fields, current = [], []
    in_quote = False
    i = 0
    while i < len(row_text):
        ch = row_text[i]
        if ch == "'":
            if in_quote and i + 1 < len(row_text) and row_text[i + 1] == "'":
                current.extend(["'", "'"])
                i += 2
                continue
            in_quote = not in_quote
            current.append(ch)
        elif ch == "," and not in_quote:
            fields.append("".join(current).strip())
            current = []
        else:
            current.append(ch)
        i += 1
    fields.append("".join(current).strip())
    return fields


def parse_sql_literal(value):
    value = value.strip()
    if value.upper() == "NULL":
        return ""
    if len(value) >= 2 and value[0] == "'" and value[-1] == "'":
        return value[1:-1].replace("''", "'")
    return value


def extract_insert_rows(sql_text):
    """Read VALUES tuples while respecting quoted text and JSON strings."""
    match = re.search(r"\bVALUES\b", sql_text, flags=re.IGNORECASE)
    if not match:
        raise ValueError("No INSERT ... VALUES section found in responses_rows.sql")

    rows = []
    in_quote = False
    depth = 0
    row_chars = []
    i = match.end()
    while i < len(sql_text):
        ch = sql_text[i]
        if ch == "'":
            if in_quote and i + 1 < len(sql_text) and sql_text[i + 1] == "'":
                if depth > 0:
                    row_chars.extend(["'", "'"])
                i += 2
                continue
            in_quote = not in_quote
            if depth > 0:
                row_chars.append(ch)
        elif not in_quote:
            if ch == "(":
                if depth == 0:
                    row_chars = []
                else:
                    row_chars.append(ch)
                depth += 1
            elif ch == ")":
                depth -= 1
                if depth == 0:
                    rows.append("".join(row_chars))
                elif depth > 0:
                    row_chars.append(ch)
                if depth < 0:
                    break
            elif depth > 0:
                row_chars.append(ch)
            elif ch == ";":
                break
        elif depth > 0:
            row_chars.append(ch)
        i += 1

    parsed = []
    for row in rows:
        fields = [parse_sql_literal(v) for v in split_sql_fields(row)]
        if len(fields) != 7:
            print(f"Warning: skipping row with {len(fields)} columns (expected 7).")
            continue
        parsed.append({
            "id": fields[0],
            "form_id": fields[1],
            "job_title": fields[2],
            "nature_of_work": fields[3],
            "evaluation_answers": fields[4],
            "text_answers": fields[5],
            "submitted_at": fields[6],
        })
    return parsed


if not SQL_PATH.exists():
    raise FileNotFoundError(
        f"Could not find {SQL_PATH.name}. Put it in the same folder as this script."
    )

sql_content = SQL_PATH.read_text(encoding="utf-8-sig")
responses = extract_insert_rows(sql_content)
if not responses:
    raise ValueError(
        "No response rows were parsed. Check the SQL INSERT statement format."
    )

mcq_stats = defaultdict(list)
for resp in responses:
    try:
        eval_answers = json.loads(resp["evaluation_answers"])
    except (json.JSONDecodeError, TypeError):
        print(f"Warning: invalid evaluation_answers JSON for response {resp['id']}")
        continue
    if isinstance(eval_answers, dict):
        for idx, ans in eval_answers.items():
            try:
                mcq_stats[int(idx)].append(ans)
            except (TypeError, ValueError):
                continue

# الأهداف
objectives_questions = arabic_headers[5:7]

# تقييم البرنامج التدريبي
training_evaluation_questions = [
    "محتوى البرنامج التدريبي",
    "درجة تطوير البرنامج لمعارفك ومهاراتك",
    "ملائمة البرنامج لعملك",
    "إمكانية تطبيق ماتعلمته على عملك"
]

# المادة التدريبية
training_material_questions = [
    "مدى ارتباط المادة التدريبية بموضوعات البرنامج",
    "مدى وضوح المادة التدريبية وتنظيمها وخلوها من الأخطاء",
    "مدى احتواء المادة التدريبية على موضوعات جديدة",
    "المادة التدريبية التي وزعت في البرنامج"
]

# تقييم المدرب
trainer_evaluation_questions = [
    "الأداء الحركي والصوتي للمدرب",
    "قدرة المدرب على توصيل المعلومات",
    "تنوع أساليب العرض من حيث الوسائل المستخدمة",
    "السيطرة على الوقت",
    "مدى تعاونه مع المتدربين",
    "قدرة المدرب على تحفيز المشاركين على التفاعل"
]

# التنسيق والإشراف العلمي
coordination_questions = [
    "حضور المنسق كافة الجلسات التدريبية",
    "استماع المنسق لملاحظات واقتراحات المشاركين",
    "مدى تلبيته لمتطلبات واحتياجات المشاركين"
]

# مكان وزمان البرنامج التدريبي
time_place_questions = [
    "كفاية عدد الساعات التدريبية للموضوعات",
    "مدى ملائمة نظام الجلسات التدريبية",
    "مدى ملائمة توقيت البرنامج"
]

# الخدمات الإدارية
administrative_services_questions = [
    "ملائمة القاعة التدريبية",
    "الظروف البيئية المحيطة بالمشارك (التهوية , الإضاءة , التكييف)"
]


def get_category(question):
    if question in objectives_questions:
        return "الأهداف"
    if question in training_evaluation_questions:
        return "تقييم البرنامج التدريبي"
    if question in training_material_questions:
        return "المادة التدريبية"
    if question in trainer_evaluation_questions:
        return "تقييم المدرب"
    if question in coordination_questions:
        return "التنسيق والإشراف العلمي"
    if question in time_place_questions:
        return "مكان وزمان البرنامج التدريبي"
    if question in administrative_services_questions:
        return "الخدمات الإدارية"
    return ""


# Create one statistics row per answer option, for the 24 scored questions.
csv_rows = []
for i, question_title in enumerate(arabic_headers[5:29]):
    answers = mcq_stats.get(i, [])
    count = Counter(answers)
    total = len(answers)
    for answer_option, repetitions in count.items():
        percentage = round(repetitions / total * 100, 1) if total else 0
        csv_rows.append([
            get_category(question_title),
            question_title,
            answer_option,
            repetitions,
            f"{percentage:g}%"
        ])

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

wb = Workbook()
ws = wb.active
ws.title = "Statistics"
ws.sheet_view.rightToLeft = True

headers = ["المحور", "السؤال", "الإجابة", "عدد الإجابات", "النسبة المئوية"]
ws.append(headers)
for row in csv_rows:
    ws.append(row)

# Style the header and make the worksheet readable.
for cell in ws[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1F4E78")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(vertical="top", wrap_text=True)

widths = {1: 28, 2: 55, 3: 20, 4: 16, 5: 18}
for col_idx, width in widths.items():
    ws.column_dimensions[get_column_letter(col_idx)].width = width
ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions

wb.save(OUTPUT_PATH)
print(f"Excel file created: {OUTPUT_PATH}")
print(f"Responses parsed: {len(responses)}")
print("Responses with evaluation answers: " + str(sum(1 for resp in responses if resp["evaluation_answers"])))
print(f"Statistics rows written: {len(csv_rows)}")
